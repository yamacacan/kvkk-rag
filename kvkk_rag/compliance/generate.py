# ${...} sablonlarini kurum profili ve envanterden doldurur.
from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Any

from ..config import settings
from ..inventory.loader import InventoryRow
from . import sections
from .profile import ALIAS, OrgProfile

PH = re.compile(r"\$\{([^}]+)\}")

# Tablo basliginda kullanilan placeholder'lar liste degil etiket olmali
BASLIK_ETIKET = {
    "veri_konusu_kisi": "VERİ KONUSU KİŞİ GRUBU",
    "veri_kategorileri": "VERİ KATEGORİSİ",
    "islenme_amaclari": "KİŞİSEL VERİ İŞLEME AMACI",
    "hukuki_sebepler": "HUKUKİ SEBEP",
}

# Bu placeholder'lar bir arada bulundugu tablo satirini envanterden cogaltir
TEKRAR_GRUBU = {"kategori", "sure", "imha"}

# Sabit yazilmis bolumleri envanterden/LLM'den uretilen sablonlar
BOLUM_URETILEN = {"politika", "veri_ihlali"}

# Profilden degil envanterden/LLM'den doldugu icin "eksik alan" sayilmaz
URETILEN_PH = frozenset(sections.VERITABANI + sections.YAPAY_ZEKA)
# Baska bir profil alanindan turetilir (qr_kod <- web_adres): doldurulacak alan degildir
TURETILEN_PH = frozenset({"qr_kod"})
# Degeri buyuk harfe cevrilen baslik placeholder'lari
BUYUK_HARF = frozenset({"KURUM", "BIRIM", "ISLEME_AMACI"})
# Faaliyet bazli belgelerde envanterden dolan (profil alani sayilmayan) placeholder'lar
FAALIYET_PH = frozenset({"BIRIM", "birim", "ISLEME_AMACI", "kisisel_veriler", "alici_gruplari", "aktarim_yonu"})

NUMPR = ("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numPr")

SABLONLAR = {
    "aydinlatma": "Aydınlatma_Metni_SablonKVKK (3).docx",
    "politika": "Kisisel_Veri_Politikasi_Sablon_Final.docx",
    "saklama_imha": "Saklama_İmha_Sablon.docx",
    "veri_isleyen": "Veri_Isleyen_Protokolu_Sablon.docx",
    "veri_ihlali": "Veri_Ihlali_Mudahale_Proseduru_Sablon.docx",
    "genel_aydinlatma": "Genel_Aydinlatma_Metni_Sablon.docx",
    "basvuru_formu": "Basvuru_Formu_Sablon.docx",
    "acik_riza": "Acik_Riza_Beyani_Sablon.docx",
}

# Aydinlatma metni ve acik riza beyani her faaliyet icin ayri duzenlenir (KVKK m.10, m.5/1);
# digerleri kurum geneli. Faaliyet belgeleri envanter degisince kuyrukta otomatik yenilenir.
FAALIYET_BAZLI = {"aydinlatma", "acik_riza"}

BELGE_ADI = {
    "aydinlatma": "Aydınlatma Metni",
    "politika": "Kişisel Veri İşleme ve Koruma Politikası",
    "saklama_imha": "Kişisel Veri Saklama ve İmha Politikası",
    "veri_isleyen": "Veri İşleyen Protokolü",
    "veri_ihlali": "Veri İhlali Müdahale Prosedürü",
    "genel_aydinlatma": "Genel Aydınlatma Metni",
    "basvuru_formu": "Kişisel Veri Sahibi Başvuru Formu",
    "acik_riza": "Açık Rıza Beyanı",
}


def _replace_in_paragraph(par, degerler: dict[str, str]) -> None:
    # Word placeholder'i birden cok run'a bolebilir; run metinleri birlestirilip
    # degistirilir, bicim ilk run'dan korunur.
    tam = "".join(r.text for r in par.runs)
    if not PH.search(tam):
        return
    yeni = PH.sub(lambda m: degerler.get(m.group(1), m.group(0)), tam)
    if yeni == tam:
        return
    if par.runs:
        par.runs[0].text = yeni
        for r in par.runs[1:]:
            r.text = ""
    else:
        par.text = yeni


def _cell_text(cell) -> str:
    return "\n".join(p.text for p in cell.paragraphs)


def _set_cell(cell, metin: str) -> None:
    ilk = cell.paragraphs[0]
    if ilk.runs:
        ilk.runs[0].text = metin
        for r in ilk.runs[1:]:
            r.text = ""
    else:
        ilk.text = metin
    for p in cell.paragraphs[1:]:
        for r in p.runs:
            r.text = ""


def _kategori_satirlari(rows: list[InventoryRow]) -> list[tuple[str, str, str]]:
    # Saklama/Imha tablosu: veri kategorisi basina saklama suresi ve imha yontemi
    gruplar: dict[str, dict[str, set[str]]] = {}
    for r in rows:
        kat = (r.veri_kategorisi or "").strip()
        if not kat:
            continue
        g = gruplar.setdefault(kat, {"sure": set(), "imha": set()})
        if r.saklama_suresi:
            g["sure"].add(r.saklama_suresi.strip())
        if r.imha_yontemi:
            g["imha"].add(r.imha_yontemi.strip())

    out = []
    for kat in sorted(gruplar):
        g = gruplar[kat]
        sure = " / ".join(sorted(g["sure"])) or "BELİRLENMEDİ"
        imha = " / ".join(sorted(g["imha"])) or "BELİRLENMEDİ"
        out.append((kat, sure, imha))
    return out


def _expand_table(table, satirlar: list[tuple[str, str, str]]) -> bool:
    import copy

    hedef = None
    for row in table.rows:
        ph = set(PH.findall("\n".join(_cell_text(c) for c in row.cells)))
        if ph and ph <= TEKRAR_GRUBU:
            hedef = row
            break
    if hedef is None:
        return False

    sablon_tr = hedef._tr
    for kat, sure, imha in satirlar:
        yeni_tr = copy.deepcopy(sablon_tr)
        sablon_tr.addprevious(yeni_tr)
        yeni = table.rows[[r._tr for r in table.rows].index(yeni_tr)]
        for cell in yeni.cells:
            metin = _cell_text(cell)
            metin = (metin.replace("${kategori}", kat)
                          .replace("${sure}", sure)
                          .replace("${imha}", imha))
            _set_cell(cell, metin)

    sablon_tr.getparent().remove(sablon_tr)
    return True


def _madde_imli(par) -> bool:
    pPr = par._p.pPr
    return pPr is not None and pPr.find(NUMPR) is not None


def _expand_list(d, degerler: dict[str, str]) -> None:
    # Word madde imli bir paragrafta cok satirli liste tek ime sikisir; paragraf
    # her oge icin klonlanir ve "• " oneki Word'un kendi imine birakilir.
    import copy
    from docx.text.paragraph import Paragraph

    for par in list(d.paragraphs):
        tam = ("".join(r.text for r in par.runs) or par.text).strip()
        m = PH.fullmatch(tam)
        if not m or not _madde_imli(par):
            continue
        deger = degerler.get(m.group(1))
        if not deger or "\n" not in deger:
            continue
        for oge in [x.lstrip("• \t").strip() for x in deger.split("\n") if x.strip()]:
            yeni_el = copy.deepcopy(par._p)
            par._p.addprevious(yeni_el)
            _set_par(Paragraph(yeni_el, par._parent), oge)
        par._p.getparent().remove(par._p)


def _qr_png(metin: str, olcek: int = 8) -> bytes | None:
    # QR kod (segno, saf Python). Kutuphane yoksa None: cagiran metne duser.
    try:
        import segno
    except ImportError:
        return None
    buf = io.BytesIO()
    segno.make(metin, error="m").save(buf, kind="png", scale=olcek, border=2)
    return buf.getvalue()


def _qr_yerlestir(d, url: str) -> None:
    # Metni tam olarak ${qr_kod} olan paragraflara web adresinin QR kodu (resim) konur;
    # adres bos ya da segno yoksa paragraf adresin kendisiyle (veya bos) doldurulur.
    from docx.shared import Cm

    for par in d.paragraphs:
        tam = ("".join(r.text for r in par.runs) or par.text).strip()
        if tam != "${qr_kod}":
            continue
        png = _qr_png(url) if url else None
        _set_par(par, "" if png else (url or ""))
        if png:
            par.add_run().add_picture(io.BytesIO(png), width=Cm(4))


def _set_par(par, metin: str) -> None:
    if par.runs:
        par.runs[0].text = metin
        for r in par.runs[1:]:
            r.text = ""
    else:
        par.text = metin


def doldur(sablon_yolu: Path, profile: OrgProfile,
           rows: list[InventoryRow] | None = None,
           ekstra: dict[str, str] | None = None) -> tuple[bytes, list[str]]:
    import docx

    d = docx.Document(str(sablon_yolu))

    # Once tekrar eden tablo satirlari (kalan placeholder'lar sonra degisir)
    if rows:
        satirlar = _kategori_satirlari(rows)
        if satirlar:
            for t in d.tables:
                _expand_table(t, satirlar)

    degerler: dict[str, str] = {}
    for ph in set(PH.findall(_belge_metni(d))):
        v = profile.value_for(ph)
        if v is not None:
            # ${KURUM} / ${BIRIM} / ${ISLEME_AMACI}: baslik kullanimi, buyuk harfle (Turkce i/ı kurali)
            degerler[ph] = _tr_buyuk(v) if ph in BUYUK_HARF else v
    # Envanterden/LLM'den uretilen bolumler profil degerlerini ezer
    degerler.update({k: v for k, v in (ekstra or {}).items() if v})

    _expand_list(d, degerler)
    _qr_yerlestir(d, (profile.web_adres or "").strip())

    for par in d.paragraphs:
        _replace_in_paragraph(par, degerler)

    for t in d.tables:
        for ri, row in enumerate(t.rows):
            for cell in row.cells:
                for par in cell.paragraphs:
                    if ri == 0:
                        # Baslik satirinda liste degil etiket yazilir
                        tam = "".join(r.text for r in par.runs) or par.text
                        yeni = PH.sub(
                            lambda m: BASLIK_ETIKET.get(m.group(1), degerler.get(m.group(1), m.group(0))),
                            tam)
                        if yeni != tam:
                            if par.runs:
                                par.runs[0].text = yeni
                                for r in par.runs[1:]:
                                    r.text = ""
                            else:
                                par.text = yeni
                    else:
                        _replace_in_paragraph(par, degerler)

    kalan = sorted(set(PH.findall(_belge_metni(d))))

    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue(), kalan


def _tr_buyuk(metin: str) -> str:
    return metin.replace("i", "İ").replace("ı", "I").upper()


def _belge_metni(d) -> str:
    parcalar = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                parcalar.append(_cell_text(c))
    return "\n".join(parcalar)


def sablon_bilgisi() -> list[dict[str, Any]]:
    import docx

    out = []
    for anahtar, dosya in SABLONLAR.items():
        yol = settings.TEMPLATES_DIR / dosya
        if not yol.exists():
            out.append({"anahtar": anahtar, "ad": BELGE_ADI[anahtar],
                        "dosya": dosya, "mevcut": False, "placeholder": []})
            continue
        d = docx.Document(str(yol))
        ph = sorted(set(PH.findall(_belge_metni(d))))
        out.append({
            "anahtar": anahtar, "ad": BELGE_ADI[anahtar], "dosya": dosya,
            "mevcut": True, "placeholder": ph,
            "alanlar": sorted({ALIAS.get(x, x) for x in ph
                               if x not in URETILEN_PH and x not in TURETILEN_PH and x not in FAALIYET_PH}),
            "uretilen_bolumler": sorted(set(ph) & URETILEN_PH),
        })
    return out


def bolumleri_uret(anahtar: str, profile: OrgProfile,
                   rows: list[InventoryRow] | None,
                   llm=None) -> tuple[dict[str, str], dict[str, str]]:
    # Sabit yazilmis bolumlerin envanterden/LLM'den uretilen karsiliklari.
    if anahtar not in BOLUM_URETILEN or not rows:
        return {}, {}
    from . import sections
    if anahtar == "veri_ihlali":
        return sections.ihlal_hesapla(profile, rows, llm)
    return sections.hesapla(profile, rows, llm)


def uret(anahtar: str, profile: OrgProfile,
         rows: list[InventoryRow] | None = None,
         llm=None) -> tuple[bytes, list[str], str, dict[str, str]]:
    if anahtar not in SABLONLAR:
        raise ValueError(f"Bilinmeyen şablon: {anahtar}")
    yol = settings.TEMPLATES_DIR / SABLONLAR[anahtar]
    if not yol.exists():
        raise FileNotFoundError(f"Şablon bulunamadı: {yol}")
    ekstra, kaynak = bolumleri_uret(anahtar, profile, rows, llm)
    icerik, kalan = doldur(yol, profile, rows, ekstra)
    ad = f"{BELGE_ADI[anahtar].replace(' ', '_')}.docx"
    return icerik, kalan, ad, kaynak
