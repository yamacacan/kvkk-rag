# Envanteri KVKK rehberine uygun, bicimlendirilmis Excel olarak disa aktarir.
from __future__ import annotations

import io
from datetime import date
from typing import Any

from . import taxonomy
from .loader import InventoryRow

# Ana sayfa: kaynak veri_envanteri.xlsx ile BIREBIR ayni 14 sutun (ad ve sira).
# Bu duzen KVKK Envanter Hazirlama Rehberi EK-1 ornegiyle alan bazinda ortusuyor;
# Sicil Yonetmeligi m.4/1-(h)'deki 7 asgari unsurun tamami burada.
SUTUNLAR: list[tuple[str, str, int]] = [
    ("satir_no",            "ID",                           7),
    ("birim",               "Birim",                       28),
    ("faaliyet",            "Faaliyet",                    30),
    ("veri_kategorisi",     "Veri Kategorisi",             22),
    ("kisisel_veri",        "Kişisel Veri",                34),
    ("ozel_nitelikli_veri", "Özel Nitelikli Kişisel Veri", 26),
    ("kisi_grubu",          "Veri Konusu Kişi Grubu",      24),
    ("isleme_amaci",        "İşleme Amacı",                34),
    ("hukuki_sebep",        "Hukuki Sebep",                40),
    ("saklama_suresi",      "Saklama Süresi",              28),
    ("alici_grubu",         "Alıcı Grupları",              26),
    ("yurt_disi_aktarim",   "Yurt Dışı Aktarım",           16),
    ("teknik_tedbir",       "Teknik Tedbirler",            46),
    ("idari_tedbir",        "İdari Tedbirler",             46),
]

# Rehberin "ilave edilebilir" dedigi tamamlayici alanlar. Kaynak dosyada sutun
# olarak yok; ana sayfanin formatini bozmamak icin ayri sayfaya yazilir.
EK_SUTUNLAR: list[tuple[str, str, int]] = [
    ("satir_no",              "ID",                     7),
    ("birim",                 "Birim",                 28),
    ("faaliyet",              "Faaliyet",              30),
    ("kisisel_veri",          "Kişisel Veri",          34),
    ("periyodik_imha_suresi", "Periyodik İmha Süresi", 20),
    ("imha_yontemi",          "İmha Yöntemi",          28),
    ("kayit_ortami",          "Kayıt Ortamı",          28),
    ("yurt_disi_ulke",        "Aktarılan Ülke",        20),
    ("aktarim_amaci",         "Aktarım Amacı",         30),
    ("veri_isleyen",          "Veri İşleyen",          26),
]

LACIVERT = "1F2A44"
BEYAZ = "FFFFFF"
GRI = "F2F4F7"
KIRMIZI = "FDE8E8"
TURUNCU = "FEF3C7"
MOR = "F5E8FF"
YESIL = "E7F8F0"


def _stiller():
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    ince = Side(style="thin", color="D0D5DD")
    return {
        "baslik_font": Font(name="Calibri", size=10, bold=True, color=BEYAZ),
        "baslik_dolgu": PatternFill("solid", fgColor=LACIVERT),
        "hucre_font": Font(name="Calibri", size=9),
        "kalin": Font(name="Calibri", size=9, bold=True),
        "hiza": Alignment(vertical="top", wrap_text=True),
        "hiza_orta": Alignment(horizontal="center", vertical="center", wrap_text=True),
        "kenar": Border(left=ince, right=ince, top=ince, bottom=ince),
        "dolgu": lambda renk: PatternFill("solid", fgColor=renk),
    }


def _satir_rengi(row: InventoryRow) -> str | None:
    sev = {b["seviye"] for b in row.bulgular}
    if "kritik" in sev:
        return KIRMIZI
    if "yuksek" in sev:
        return TURUNCU
    if sev:
        return MOR
    return None


def build(rows: list[InventoryRow], kurum: str = "", ozet: dict[str, Any] | None = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    st = _stiller()
    wb = Workbook()
    ozel_set = taxonomy.OZEL_NITELIKLI_KATEGORILER

    def _tablo(ws, sutunlar, satirlar, renkli: bool):
        # Baslik 1. satirda, veri 2. satirdan: kaynak dosyayla ayni yerlesim.
        # Bos hucre bos kalir ("—" yazilmaz); dosya sisteme geri yuklenebilsin.
        for i, (_, etiket, genislik) in enumerate(sutunlar, start=1):
            h = ws.cell(row=1, column=i, value=etiket)
            h.font = st["baslik_font"]
            h.fill = st["baslik_dolgu"]
            h.alignment = st["hiza_orta"]
            h.border = st["kenar"]
            ws.column_dimensions[get_column_letter(i)].width = genislik
        ws.row_dimensions[1].height = 30
        for r, row in enumerate(satirlar, start=2):
            renk = _satir_rengi(row) if renkli else None
            for i, (alan, _, _) in enumerate(sutunlar, start=1):
                deger = getattr(row, alan, None)
                c = ws.cell(row=r, column=i, value=deger if deger not in (None, "") else None)
                c.font = st["hucre_font"]
                c.alignment = st["hiza"]
                c.border = st["kenar"]
                if renk:
                    c.fill = st["dolgu"](renk)
                if alan == "veri_kategorisi" and deger and deger in ozel_set:
                    c.font = st["kalin"]
                    c.fill = st["dolgu"](MOR)
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(sutunlar))}{max(1, len(satirlar) + 1)}"
        ws.sheet_view.zoomScale = 90

    # --- Sayfa 1: Envanter (kaynak dosya formati) ---
    ws = wb.active
    ws.title = "Veri İşleme Envanteri"
    _tablo(ws, SUTUNLAR, rows, renkli=True)

    # --- Sayfa 2: Ek alanlar (yalnizca dolu olan satirlar) ---
    ek_alanlar = [a for a, _, _ in EK_SUTUNLAR if a not in ("satir_no", "birim", "faaliyet", "kisisel_veri")]
    ek_satirlar = [r for r in rows if any(getattr(r, a, None) for a in ek_alanlar)]
    if ek_satirlar:
        ws_ek = wb.create_sheet("Ek Alanlar")
        _tablo(ws_ek, EK_SUTUNLAR, ek_satirlar, renkli=False)

    # --- Sayfa 3: Satir bazli uyum bulgulari ---
    bulgulu = [(r, b) for r in rows for b in r.bulgular]
    if bulgulu:
        ws_b = wb.create_sheet("Uyum Bulguları")
        basliklar = [("ID", 7), ("Birim", 28), ("Faaliyet", 30), ("Kişisel Veri", 30),
                     ("Seviye", 10), ("Kod", 10), ("Bulgu", 46), ("Mevzuat Dayanağı", 36)]
        for i, (etiket, genislik) in enumerate(basliklar, start=1):
            h = ws_b.cell(row=1, column=i, value=etiket)
            h.font = st["baslik_font"]; h.fill = st["baslik_dolgu"]
            h.alignment = st["hiza_orta"]; h.border = st["kenar"]
            ws_b.column_dimensions[get_column_letter(i)].width = genislik
        for i, (r, b) in enumerate(bulgulu, start=2):
            renk = {"kritik": KIRMIZI, "yuksek": TURUNCU}.get(b["seviye"], MOR)
            for j, v in enumerate((r.satir_no, r.birim, r.faaliyet, r.kisisel_veri,
                                   b["seviye"].upper(), b["kod"], b["baslik"], b["dayanak"]), start=1):
                c = ws_b.cell(row=i, column=j, value=v)
                c.font = st["hucre_font"]; c.alignment = st["hiza"]
                c.border = st["kenar"]; c.fill = st["dolgu"](renk)
        ws_b.freeze_panes = "A2"
        ws_b.auto_filter.ref = f"A1:H{len(bulgulu) + 1}"

    # --- Sayfa 4: Uyum Özeti ---
    ws2 = wb.create_sheet("Uyum Özeti")
    ws2.column_dimensions["A"].width = 14
    ws2.column_dimensions["B"].width = 52
    ws2.column_dimensions["C"].width = 40
    ws2.column_dimensions["D"].width = 12

    ws2.merge_cells("A1:D1")
    t = ws2.cell(row=1, column=1, value="Uyum Denetimi Özeti")
    t.font = st["baslik_font"]
    t.fill = st["baslik_dolgu"]
    t.alignment = st["hiza_orta"]
    ws2.row_dimensions[1].height = 24

    o = ozet or {}
    bilgi = [
        ("Kurum", kurum or "—"),
        ("Düzenleme tarihi", f"{date.today():%d.%m.%Y}"),
        ("Dayanak", "6698 sayılı KVKK · Veri Sorumluları Sicili Hk. Yönetmelik m.4/1-(h) · "
                    "KVKK Kişisel Veri İşleme Envanteri Hazırlama Rehberi"),
        ("Toplam kayıt", len(rows)),
        ("Bulgu sayısı", sum(len(r.bulgular) for r in rows)),
        ("Bulgusuz kayıt", sum(1 for r in rows if not r.bulgular)),
        ("Uyum oranı", f"%{(sum(1 for r in rows if not r.bulgular) / len(rows) * 100):.1f}" if rows else "—"),
        ("Özel nitelikli veri içeren kayıt",
         sum(1 for r in rows if r.veri_kategorisi in ozel_set or r.ozel_nitelikli_veri)),
        ("Yurt dışına aktarım yapılan kayıt",
         sum(1 for r in rows if (r.yurt_disi_aktarim or "").strip().lower() in ("evet", "var"))),
    ]
    for i, (k, v) in enumerate(bilgi, start=3):
        ws2.cell(row=i, column=1, value=k).font = st["kalin"]
        ws2.cell(row=i, column=2, value=v).font = st["hucre_font"]

    sat = len(bilgi) + 5
    for i, etiket in enumerate(("Kod", "Bulgu", "Mevzuat Dayanağı", "Adet"), start=1):
        h = ws2.cell(row=sat, column=i, value=etiket)
        h.font = st["baslik_font"]
        h.fill = st["baslik_dolgu"]
        h.alignment = st["hiza_orta"]

    kodlar: dict[str, dict[str, Any]] = {}
    for r in rows:
        for b in r.bulgular:
            k = kodlar.setdefault(b["kod"], {"baslik": b["baslik"], "dayanak": b["dayanak"],
                                             "adet": 0, "seviye": b["seviye"]})
            k["adet"] += 1
    for i, (kod, d) in enumerate(sorted(kodlar.items()), start=sat + 1):
        renk = {"kritik": KIRMIZI, "yuksek": TURUNCU}.get(d["seviye"], MOR)
        for j, v in enumerate((kod, d["baslik"], d["dayanak"], d["adet"]), start=1):
            c = ws2.cell(row=i, column=j, value=v)
            c.font = st["hucre_font"]
            c.alignment = st["hiza"]
            c.border = st["kenar"]
            c.fill = st["dolgu"](renk)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
