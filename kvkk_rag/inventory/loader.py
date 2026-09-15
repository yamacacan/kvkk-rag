# veri_envanteri.xlsx okur ve satirlari kanonik alan adlariyla dondurur.
from __future__ import annotations

from dataclasses import dataclass, field, asdict
from pathlib import Path
from typing import Any

from ..config import settings

COLUMNS = {
    "ID": "id",
    "Birim": "birim",
    "Faaliyet": "faaliyet",
    "Veri Kategorisi": "veri_kategorisi",
    "Kişisel Veri": "kisisel_veri",
    "Özel Nitelikli Kişisel Veri": "ozel_nitelikli_veri",
    "Veri Konusu Kişi Grubu": "kisi_grubu",
    "İşleme Amacı": "isleme_amaci",
    "Hukuki Sebep": "hukuki_sebep",
    "Saklama Süresi": "saklama_suresi",
    "Alıcı Grupları": "alici_grubu",
    "Yurt Dışı Aktarım": "yurt_disi_aktarim",
    "Teknik Tedbirler": "teknik_tedbir",
    "İdari Tedbirler": "idari_tedbir",
}

# Sicil Yonetmeligi m.4/1-(h) envanterin zorunlu icerigini sayar. Asagidakiler
# xlsx'te sutun olarak yok; rehberin istedigi tamamlayici alanlardir.
TURETILEN_ALANLAR = ("imha_yontemi", "kayit_ortami", "periyodik_imha_suresi",
                     "yurt_disi_ulke", "aktarim_amaci", "veri_isleyen")


@dataclass
class InventoryRow:
    satir_no: int
    id: str | None = None
    birim: str | None = None
    faaliyet: str | None = None
    veri_kategorisi: str | None = None
    kisisel_veri: str | None = None
    ozel_nitelikli_veri: str | None = None
    kisi_grubu: str | None = None
    isleme_amaci: str | None = None
    hukuki_sebep: str | None = None
    saklama_suresi: str | None = None
    alici_grubu: str | None = None
    yurt_disi_aktarim: str | None = None
    teknik_tedbir: str | None = None
    idari_tedbir: str | None = None
    imha_yontemi: str | None = None
    kayit_ortami: str | None = None
    periyodik_imha_suresi: str | None = None
    yurt_disi_ulke: str | None = None
    aktarim_amaci: str | None = None
    veri_isleyen: str | None = None
    bulgular: list[dict[str, Any]] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


def _clean(v: Any) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def load(path: Path | None = None) -> tuple[list[InventoryRow], list[str]]:
    import openpyxl

    path = path or settings.INVENTORY_PATH
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]

    rows = ws.iter_rows(values_only=True)
    header = [_clean(h) or "" for h in next(rows)]
    idx = {COLUMNS[h]: i for i, h in enumerate(header) if h in COLUMNS}
    eksik = [h for h in COLUMNS if h not in header]

    out: list[InventoryRow] = []
    for n, raw in enumerate(rows, start=2):
        if all(c is None for c in raw):
            continue
        kwargs = {alan: _clean(raw[i]) for alan, i in idx.items() if i < len(raw)}
        out.append(InventoryRow(satir_no=n, **kwargs))

    wb.close()
    return out, eksik


def unique_values(rows: list[InventoryRow], field_name: str) -> list[str]:
    return sorted({v for v in (getattr(r, field_name) for r in rows) if v})
