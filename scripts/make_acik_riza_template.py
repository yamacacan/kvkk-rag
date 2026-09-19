"""Açık Rıza Beyanı şablonunu (docx) üretir — faaliyet bazlı: birim, işleme amacı, kişisel
veriler, alıcı grupları ve aktarım yönü envanterdeki faaliyetten dolar. Envanter değişince
Faaliyet Belgeleri kuyruğu bu belgeyi otomatik yeniler.

Placeholder'lar: ${KURUM}, ${kurum}, ${BIRIM}, ${ISLEME_AMACI}, ${kisisel_veriler},
${alici_gruplari}, ${aktarim_yonu}.

    python scripts/make_acik_riza_template.py
"""
from __future__ import annotations

import sys
from pathlib import Path

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from make_basvuru_formu_template import _hucre_yaz  # noqa: E402
from make_veri_ihlali_template import KAYNAK, _prototipler, _yeni_paragraf  # noqa: E402

HEDEF = Path(__file__).resolve().parents[1] / "data" / "templates" / "Acik_Riza_Beyani_Sablon.docx"

ICERIK: list[tuple[str, str]] = [
    ("orta_kalin", "${KURUM}"),
    ("h1_orta", "KİŞİSEL VERİLERİN İŞLENMESİ"),
    ("h1_orta", "AÇIK RIZA BEYANI (${BIRIM} İŞLEMLERİ)"),
    ("p", "${kurum} (\"VERİ SORUMLUSU\") ve gerekli güvenlik önlemlerinin alınması suretiyle yetkilendirdiği veri "
          "işleyenler tarafından, 6698 Sayılı Kişisel Verilerin Korunması Kanunu (\"KANUN\") ve bu KANUN ile ilgili "
          "mevzuat hükümlerine uygun olarak bilgime sunulan \"Kişisel Verilerin Korunması Aydınlatma Metni\" ve "
          "\"Kişisel Verilerin Korunması ve İşlenmesi Politikası\" çerçevesinde,"),
    ("p_kalin", "VERİ SORUMLUSU tarafından ${ISLEME_AMACI} amacıyla"),
    ("p", "Kişisel verilerimin (${kisisel_veriler}) gerekli olan süre kadar muhafaza edilme ilkesi başta olmak üzere "
          "KANUN'un 4. maddesinde ifade edilen genel ilkelere uygun şekilde işlenebileceğini ve ${alici_gruplari} gibi "
          "mercilere başta olmak üzere ${aktarim_yonu} aktarılabileceğini, VERİ SORUMLUSU ile paylaşmış olduğum kişisel "
          "verilerin doğru ve güncel olduğunu; işbu bilgilerde değişiklik olması halinde değişiklikleri VERİ "
          "SORUMLUSU'na bildireceğimi ve bu hususta VERİ SORUMLUSU tarafından şahsıma gerekli aydınlatmanın ve "
          "bilgilendirmenin yapıldığını, açık rıza beyanımın tarafımdan yazılı olarak feshedilmediği sürece geçerli "
          "olduğunu, işbu \"Açık Rıza Beyanı\"nı, bilgilendirmeye dayalı olarak \"Kişisel Verilerin Korunması Aydınlatma "
          "Metni\"ni okuduğumu ve anladığımı, işbu açık rızamın sonuçları üzerinde tam bilgi sahibi olduğumu, hiçbir "
          "baskı ve tehdit altında kalmadan özgür irademle, açık rıza beyanını"),
    ("kutu", "Kabul ediyorum"),
    ("kutu", "Kabul etmiyorum"),
    ("bos", ""),
]


def uret() -> Path:
    d = docx.Document(str(KAYNAK))
    proto = _prototipler(d)
    body = d.element.body
    sect = body.sectPr
    for el in list(body):
        if el is not sect:
            body.remove(el)

    def par(tur: str, metin: str):
        p = _yeni_paragraf(body, proto[tur], metin)
        if sect is not None:
            sect.addprevious(p._p)
        return p

    for tur, metin in ICERIK:
        if tur == "bos":
            par("p", "")
        elif tur == "h1_orta":
            par("h1", metin).alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif tur == "orta_kalin":
            p = par("p", metin); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.bold = True
        elif tur == "p_kalin":
            p = par("p", metin)
            for r in p.runs:
                r.font.bold = True
        elif tur == "kutu":
            par("p", "☐  " + metin).paragraph_format.space_after = Pt(2)
        else:
            par(tur, metin)

    # Imza blogu: Veri sahibinin adi soyadi / tarih / imza
    t = d.add_table(rows=3, cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for ri, etiket in enumerate(("Veri Sahibinin Adı ve Soyadı", "Tarih", "İmza")):
        t.cell(ri, 0).width = Cm(6); t.cell(ri, 1).width = Cm(10.9)
        _hucre_yaz(t.cell(ri, 0), etiket, kalin=True, boyut=11)
        _hucre_yaz(t.cell(ri, 1), "\n" if etiket == "İmza" else "", boyut=11)

    d.core_properties.title = "Açık Rıza Beyanı"
    d.save(str(HEDEF))
    return HEDEF


if __name__ == "__main__":
    print(f"Şablon yazıldı: {uret()}")
