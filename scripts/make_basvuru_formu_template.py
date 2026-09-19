"""Kişisel Veri Sahibi Başvuru Formu şablonunu (docx) üretir — ilgili kişinin (KVKK m.11/13)
doldurup ileteceği boş form: bilgi tablosu, başvuru yöntemleri tablosu, kimlik alanları,
onay kutuları, talep alanı, cevap yöntemi, açıklama ve imza bloğu. Altbilgide "Sayfa X / Y".

Biçim mevcut şablonlardan alınır (make_veri_ihlali_template ile aynı prototipler).
Placeholder'lar: ${KURUM} (başlık), ${kurum}, ${adres}, ${tarih}.

    python scripts/make_basvuru_formu_template.py
"""
from __future__ import annotations

import sys
from pathlib import Path

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from make_veri_ihlali_template import KAYNAK, _prototipler, _yeni_paragraf  # noqa: E402

HEDEF = Path(__file__).resolve().parents[1] / "data" / "templates" / "Basvuru_Formu_Sablon.docx"
FONT = "Times New Roman"
ZARF = "Zarfın üzerine \"Kişisel Verilerin Korunması Kanunu Kapsamında Bilgi Talebi\" yazılmalıdır."
ADRES_HUCRE = ["${kurum}", "${adres}"]

# (tur, veri) — tur: orta_kalin | h1_orta | h1 | p | kalin | kutu | tablo | bos | orta_kucuk
ICERIK: list[tuple[str, object]] = [
    ("orta_kalin", "T.C."),
    ("orta_kalin", "${KURUM}"),
    ("h1_orta", "KİŞİSEL VERİ SAHİBİ BAŞVURU FORMU"),
    ("tablo", {"sutun": [4.5, 12.4], "satirlar": [
        ["Doküman İçeriği", "Bu form, KVKK'nın 11. maddesi kapsamında kişisel veri sahibi haklarının kullanılması amacıyla hazırlanmıştır."],
        ["Versiyon No", "1.0"],
        ["Dayanak", "6698 sayılı Kişisel Verilerin Korunması Kanunu"],
        ["Hazırlayan", "${kurum}"],
    ], "ilk_sutun_kalin": True}),
    ("h1", "1. BAŞVURU YÖNTEMİ"),
    ("p", "Kişisel Verilerin Korunması Kanunu'nun 11. maddesinde sayılan haklarınız kapsamındaki taleplerinizi, Kanun'un 13. "
          "maddesi ile Veri Sorumlusuna Başvuru Usul ve Esasları Hakkında Tebliğ'in 5. maddesi gereğince, işbu form ile "
          "aşağıda açıklanan yöntemlerden biriyle Kurumumuza iletebilirsiniz."),
    ("tablo", {"sutun": [3.2, 5.4, 4.3, 4.0], "baslik": True, "satirlar": [
        ["BAŞVURU YÖNTEMİ", "BAŞVURUDA GEREKENLER", "BAŞVURU ADRESİ", "DİĞER İSTENENLER"],
        ["Şahsen Başvuru",
         "VERİ SORUMLUSU'nun faaliyet gösterdiği adrese kimliğinizi doğrulayarak şahsen veya vekaletname ibraz etmek "
         "suretiyle bir vekil aracılığıyla başvuruda bulunabilirsiniz.",
         ADRES_HUCRE, "Başvuru kapalı zarf ile yapılmalı, " + ZARF[0].lower() + ZARF[1:]],
        ["Posta Yoluyla Başvuru",
         "Islak imzalı başvuru formu veya dilekçe posta yoluyla gönderilerek de başvuruda bulunulabilir.",
         ADRES_HUCRE, ZARF],
        ["Noter Yoluyla Başvuru",
         "Bizzat veya vekil aracılığıyla noter kanalıyla başvuru da yapılabilir.",
         ADRES_HUCRE, "Cevabın hangi yöntemle alınmak istendiği belirtilmelidir."],
    ]}),
    ("p_onemli", "Kurumumuz talebinizi, niteliklerine göre en kısa sürede ve en geç otuz gün içinde ücretsiz olarak "
                 "sonuçlandırır. Ancak, işlemin ayrıca bir maliyet gerektirmesi halinde, Kişisel Verileri Koruma Kurulu "
                 "tarafından belirlenen tarifedeki ücretleri tarafınızdan talep edebilecektir."),
    ("h1", "2. KİMLİK VE İLETİŞİM BİLGİLERİ"),
    ("p", "Lütfen Kurumumuz tarafından sizinle iletişime geçilebilmesi için aşağıdaki alanları doldurunuz."),
    ("tablo", {"sutun": [4.5, 12.4], "satirlar": [["Adı Soyadı:", ""], ["T.C. Kimlik No:", ""], ["Telefon No:", ""],
                                                  ["E-Posta Adresi:", ""], ["Adres:", "\n"]], "ilk_sutun_kalin": True}),
    ("h1", "3. KURUMUMUZ İLE OLAN İLİŞKİNİZ"),
    ("p", "Lütfen Kurumumuz ile olan ilişkinizi (çalışan, eski çalışan, çalışan adayı, ziyaretçi, üçüncü taraf firma "
          "çalışanı ya da temsilcisi/yetkilisi gibi) belirtiniz."),
    ("kutu", "Vatandaş"), ("kutu", "Çalışan"), ("kutu", "Ziyaretçi"), ("kutu", "Diğer: ......................................"),
    ("h1", "4. KİŞİSEL VERİLERİN KORUNMASI KANUNU KAPSAMINDAKİ TALEPLERİNİZ"),
    ("p", "Lütfen Kişisel Verilerin Korunması Kanunu'nun 11. maddesi kapsamında sahip olduğunuz haklardan hangilerini "
          "kullanmak istediğinizi ve taleplerinizi ayrıntılı olarak belirtiniz."),
    ("kalin", "Talebinizin Ayrıntıları:"),
    ("tablo", {"sutun": [16.9], "satirlar": [["\n" * 12]]}),
    ("h1", "5. KURUMUMUZUN CEVABININ TARAFINIZA BİLDİRİLME YÖNTEMİ"),
    ("p", "Lütfen Kurumumuz tarafından başvurunuza verilecek cevabın tarafınıza bildirilme yöntemini seçiniz."),
    ("kutu", "E-Posta Adresime Gönderilmesini İstiyorum"), ("kutu", "Elden Bizzat Teslim Almak İstiyorum"),
    ("kutu", "Adresime Posta Yoluyla Gönderilmesini İstiyorum"), ("kutu", "Faks Yoluyla Gönderilmesini İstiyorum"),
    ("kalin", "Başvurunun Gönderileceği Adres:"),
    ("tablo", {"sutun": [16.9], "satirlar": [["\n\n"]]}),
    ("h1", "6. AÇIKLAMA"),
    ("p", "İşbu başvuru formu, Kurumumuz ile olan ilişkinizin tespiti, varsa, Kurumumuz tarafından işlenen kişisel "
          "verilerinizin eksiksiz olarak belirlenmesi ve söz konusu başvurunuza doğru ve yasal süre içerisinde cevap "
          "verilebilmesi için düzenlenmiştir."),
    ("p", "Hukuka ve kanuna aykırı bir şekilde veri paylaşımından kaynaklanabilecek hukuki risklerin ortadan kaldırılması "
          "ve kişisel verilerinizin güvenliğinin sağlanması amacına yönelik olarak Kurumumuz kimlik tespiti yapabilmek "
          "amacıyla kimlik belgesi talep etme hakkını saklı tutar."),
    ("p", "İşbu form ile ilettiğiniz taleplerinize ilişkin bilgilerin doğru veya güncel olmaması ya da başvurunun yetkisiz "
          "kişi/kişiler tarafından yapılması halinde Kurumumuz, söz konusu taleplerden dolayı sorumluluk kabul etmemekte "
          "olup, başvuruyu cevaplamak zorunda değildir."),
    ("p", "Kurumumuz, gerekli gördüğü hallerde, başvuru sahibinin kimliğini doğrulamak amacıyla ek bilgi ve belge talep "
          "edebilir."),
    ("p", "Başvuru formu eksiksiz doldurulmalıdır. Eksik bilgi nedeniyle başvuru değerlendirilemezse, başvuru sahibi "
          "bilgilendirilecek ve eksik bilgilerin tamamlanması istenecektir."),
    ("bos", ""),
    ("tablo", {"sutun": [5.5, 11.4], "satirlar": [["Başvuru Sahibi (Ad Soyad):", ""], ["Başvuru Tarihi:", "....... / ....... / .............."],
                                                  ["İmza:", "\n\n"]], "ilk_sutun_kalin": True}),
    ("bos", ""),
    ("orta_kucuk", "Bu form ${tarih} tarihinde ${kurum} tarafından hazırlanmıştır."),
    ("orta_kucuk", "Kişisel Verilerin Korunması Kanunu kapsamında hazırlanmıştır."),
]


def _hucre_yaz(cell, metin: str | list[str], kalin: bool = False, boyut: int = 10) -> None:
    satirlar = metin if isinstance(metin, list) else metin.split("\n")
    ilk = cell.paragraphs[0]
    for i, s in enumerate(satirlar):
        par = ilk if i == 0 else cell.add_paragraph()
        par.paragraph_format.space_after = Pt(2)
        par.paragraph_format.space_before = Pt(2)
        r = par.add_run(s)
        r.font.name = FONT
        r.font.size = Pt(boyut)
        r.font.bold = kalin
        r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def _tablo(d, body, spec: dict) -> None:
    satirlar = spec["satirlar"]
    t = d.add_table(rows=len(satirlar), cols=len(spec["sutun"]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for ri, satir in enumerate(satirlar):
        for ci, metin in enumerate(satir):
            hucre = t.cell(ri, ci)
            hucre.width = Cm(spec["sutun"][ci])
            baslik = spec.get("baslik") and ri == 0
            kalin = bool(baslik or (spec.get("ilk_sutun_kalin") and ci == 0))
            _hucre_yaz(hucre, metin, kalin=kalin, boyut=10)
            if baslik:
                # koyu lacivert baslik zemini, beyaz yazi (mevcut sablonlarin tablo basligi gibi)
                tcPr = hucre._tc.get_or_add_tcPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "1F2A44")
                tcPr.append(shd)
                for p in hucre.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.color.rgb = docx.shared.RGBColor(0xFF, 0xFF, 0xFF)
    # add_table tabloyu sectPr'den hemen once ekler; ardina bosluk paragrafi
    _par(body, PROTO["p"], "")


def _alan(par, alan: str) -> None:
    # Word alan kodu (PAGE / NUMPAGES) — altbilgi sayfa numarasi
    r = par.add_run()
    f = OxmlElement("w:fldSimple"); f.set(qn("w:instr"), alan)
    rr = OxmlElement("w:r"); tt = OxmlElement("w:t"); tt.text = "1"; rr.append(tt); f.append(rr)
    r._r.append(f)


def _altbilgi(d) -> None:
    sec = d.sections[0]
    par = sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in par.runs:
        r.text = ""
    r = par.add_run("Sayfa "); r.font.size = Pt(9); r.font.name = FONT
    _alan(par, "PAGE")
    r = par.add_run(" / "); r.font.size = Pt(9); r.font.name = FONT
    _alan(par, "NUMPAGES")


PROTO: dict = {}


def _par(body, proto, metin: str):
    # Paragrafi sectPr'den once ekler (sectPr yerinde kalir; add_table de ona gore konumlanir)
    p = _yeni_paragraf(body, proto, metin)
    sect = body.sectPr
    if sect is not None:
        sect.addprevious(p._p)
    return p


def uret() -> Path:
    global PROTO
    d = docx.Document(str(KAYNAK))
    PROTO = _prototipler(d)
    body = d.element.body
    sect = body.sectPr
    for el in list(body):
        if el is not sect:
            body.remove(el)

    for tur, veri in ICERIK:
        if tur == "tablo":
            _tablo(d, body, veri)
        elif tur == "bos":
            _par(body, PROTO["p"], "")
        elif tur == "h1_orta":
            _par(body, PROTO["h1"], veri).alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif tur == "h1":
            _par(body, PROTO["h1"], veri)
        elif tur == "orta_kalin":
            p = _par(body, PROTO["p"], veri); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.bold = True
        elif tur == "orta_kucuk":
            p = _par(body, PROTO["p"], veri); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.size = Pt(10); r.font.italic = True
        elif tur == "kalin":
            p = _par(body, PROTO["p"], veri)
            for r in p.runs:
                r.font.bold = True
        elif tur == "kutu":
            _par(body, PROTO["p"], "☐  " + veri).paragraph_format.space_after = Pt(2)
        elif tur == "p_onemli":
            p = _par(body, PROTO["p"], "Önemli: ")
            p.runs[0].font.bold = True
            r = p.add_run(veri); r.font.name = FONT; r.font.size = Pt(12)
        else:
            _par(body, PROTO[tur], veri)

    _altbilgi(d)
    d.core_properties.title = "Kişisel Veri Sahibi Başvuru Formu"
    d.save(str(HEDEF))
    return HEDEF


if __name__ == "__main__":
    print(f"Şablon yazıldı: {uret()}")
