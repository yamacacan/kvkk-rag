"""Veri İhlali Müdahale Prosedürü şablonunu (docx) üretir.

Biçim, mevcut Saklama/İmha şablonundan alınır (aynı başlık, gövde ve madde imi
biçimleri kopyalanır); kurum adı ${kurum}, tarih ${tarih} placeholder'ıdır. Envanterden
ve yapay zekâdan üretilen bölümler ${ihlal_*} placeholder'larıyla işaretlidir
(kvkk_rag/compliance/sections.py: ihlal_hesapla).

    python scripts/make_veri_ihlali_template.py
"""
from __future__ import annotations

import copy
import sys
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[1]
KAYNAK = ROOT / "data" / "templates" / "Saklama_İmha_Sablon.docx"
HEDEF = ROOT / "data" / "templates" / "Veri_Ihlali_Mudahale_Proseduru_Sablon.docx"

# (tur, metin) — tur: baslik | h1 | h2 | p | madde | bos | sayfa
ICERIK: list[tuple[str, str]] = [
    ("baslik", "${KURUM}"),
    ("altbaslik", "VERİ İHLALİ MÜDAHALE PROSEDÜRÜ"),
    ("orta", "Versiyon: 1.0"),
    ("orta", "Tarih: ${tarih}"),
    ("sayfa", ""),
    ("h1", "1. AMAÇ"),
    ("p", "${kurum} (\"VERİ SORUMLUSU\"), 6698 sayılı Kişisel Verilerin Korunması Kanunu'nun (\"KANUN\") 12. maddesinin "
          "5. fıkrasına göre işlenen kişisel verilerin kanuni olmayan yollarla başkaları tarafından elde edilmesi halinde, "
          "diğer bir deyişle kişisel veri ihlalinin gerçekleşmesi halinde, bu durumu en kısa sürede ilgilisine ve Kişisel "
          "Verileri Koruma Kuruluna (\"KURUL\") bildirmekle yükümlüdür."),
    ("p", "İşbu Veri İhlali Müdahale Prosedürü (\"PROSEDÜR\"), VERİ SORUMLUSU'nun işlemekte olduğu kişisel veriler ile ilgili "
          "olarak bir kişisel veri ihlalinin gerçekleşmesi halinde oluşacak krize nasıl müdahale edileceği ve atılacak "
          "adımların neler olduğu konusunda çalışanları bilgilendirmek amacıyla hazırlanmıştır."),
    ("p", "İşbu PROSEDÜR, VERİ SORUMLUSU'nun kişisel verilerin korunması ve işlemesine ilişkin yürürlüğe koymuş olduğu tüm "
          "politika ve prosedürler ile birlikte ele alınır."),
    ("h1", "2. KİŞİSEL VERİ İHLALİ"),
    ("p", "Kişisel veri ihlali, kişisel verilerin kanuna aykırı bir şekilde elde edilmesi, kişisel verilere yetkisiz erişim "
          "sağlanması, kişisel verilerin yanlışlıkla veya kasten yetkisiz kişilere açıklanması, kişisel verilerin hukuka "
          "aykırı bir şekilde silinmesi, değiştirilmesi veya bütünlüğünün bozulması gibi durumlarda ortaya çıkmaktadır."),
    ("p", "Aşağıda yer alan ve benzeri durumlar genel olarak kişisel veri ihlali olarak değerlendirilir:"),
    ("madde", "Kişisel veri içeren fiziki dokümanların veya elektronik cihazların çalınması veya kaybolması,"),
    ("madde", "Kişiye özel kullanıcı adı ve parolaların yetkisiz kişilerce ele geçirilmesi,"),
    ("madde", "Gizli bilgilerin hukuka aykırı şekilde ifşa edilmesi,"),
    ("madde", "Kişisel veri ve/veya gizli bilgi içeren e-postaların yanlışlıkla veya kasten Kurum dışında ilgisiz kişilere "
              "iletilmesi ve/veya gönderilmesi,"),
    ("madde", "Bilgi işlem ekipmanlarına, sistemlerine ve ağlarına virüs veya diğer saldırıların (örneğin siber saldırı) "
              "gerçekleşmesi suretiyle kişisel verilere hukuka aykırı erişim sağlanması."),
    ("p", "${ihlal_senaryolari}"),
    ("h1", "3. VERİ İHLALİ MÜDAHALE EKİBİ"),
    ("p", "Kişisel veri ihlali durumunda gerekli müdahalelerde bulunmak ve KANUN kapsamında öngörülen yükümlülükleri yerine "
          "getirmek için VERİ SORUMLUSU'nun bünyesinde aşağıdaki katılımcıların dahil edileceği bir Veri İhlali Müdahale "
          "Ekibi oluşturulur:"),
    ("madde", "Veri Sorumlusu İrtibat Kişisi,"),
    ("madde", "İhlalin Ortaya Çıktığı Departman Yetkilisi,"),
    ("madde", "Kişisel Verilerin Korunması Komitesi."),
    ("p", "VERİ SORUMLUSU'nun kişisel veri işleme envanterinde kayıtlı olan ve ihlal halinde departman yetkilisinin ekibe "
          "katılacağı birimler aşağıda listelenmiştir:"),
    ("madde", "${ihlal_birimleri}"),
    ("h1", "4. VERİ İHLALİ MÜDAHALE SÜRECİ"),
    ("p", "Bir veri ihlalinin gerçekleşmesi halinde, VERİ SORUMLUSU'nun bünyesinde sırasıyla veri ihlaline ilişkin bir ön "
          "değerlendirme yapılır, engelleme ve kurtarma çalışmaları yerine getirilir, ilgili riskler değerlendirilir, "
          "KURUL'a bildirim yapılır ve bunu takiben veri ihlali sonrası değerlendirme ve geliştirme yönünde neler "
          "yapılabileceği belirlenir."),
    ("h2", "4.1. VERİ İHLALİNE İLİŞKİN ÖN DEĞERLENDİRME YAPILMASI"),
    ("p", "VERİ SORUMLUSU'nun bünyesinde bir veri ihlalinin gerçekleşmesi ya da veri ihlali ihtimalinin ortaya çıkması "
          "halinde, ilgili tüm çalışanlar ilgili departmanın yetkilisine ve Veri Sorumlusu İrtibat Kişisi'ne derhal ve "
          "gecikmeksizin durumu bildirmekle yükümlüdür. Bahse konu bildirim; ihlalin gerçekleşme tarihi ve saati, ihlalin "
          "tespiti tarihi ve saati, veri ihlalinin niteliği veya veri ihlalinden etkilenenleri içerir."),
    ("p", "Veri Sorumlusu İrtibat Kişisi, Veri İhlali Müdahale Ekibi ile birlikte veri ihlaline ilişkin bildirim üzerine "
          "bir ön değerlendirme yapar."),
    ("h2", "4.2. ENGELLEME VE KURTARMA ÇALIŞMALARININ YÜRÜTÜLMESİ"),
    ("p", "Veri İhlali Müdahale Ekibi, gerçekleşen veya muhtemel veri ihlalinin VERİ SORUMLUSU'nun ve ilgili kişiler "
          "üzerindeki etkilerinin azaltılabilmesi için engelleme ve kurtarma çalışmalarını yürütür. Bu çerçevede, öncelikle "
          "veri ihlalinden haberdar edilmesi gereken departmanlar tespit edilerek ihlalin kontrol edilebilmesi, mümkünse "
          "engellenebilmesi ve zararların azaltılabilmesi için atılması gereken adımlar konusunda ilgili birimle koordineli "
          "bir şekilde bahse konu çalışmalar yürütülür. Bunu takiben, veri ihlalinden etkilenen kişi ve varsa kurum ve "
          "kuruluşlar ve iletişim bilgileri tespit edilmeye çalışılır."),
    ("h2", "4.3. RİSKLERİN TESPİT EDİLMESİ"),
    ("p", "Gerçekleşen veri ihlalinin düzeyinin belirlenmesinde ilgili kişiler üzerinde ne kadar bir potansiyel etkiye "
          "neden olduğu değerlendirilmektedir. Veri İhlali Müdahale Ekibi tarafından gerçekleşen veya muhtemel veri "
          "ihlalinin düzeyi, ihlalden etkilenen kişiler üzerinde oluşturabileceği olumsuz etkiler, ilgili riskler ortaya "
          "konularak ve değerlendirilerek tespit edilir. Söz konusu tespit işleminin yapılması sırasında; ihlalin niteliği, "
          "ihlalin nedeni, ihlale maruz kalan verinin türü, ihlalin etkisinin azaltılmasında alınan önlemler ile ihlalden "
          "etkilenen ilgili kişi kategorileri göz önünde bulundurulur."),
    ("p", "${ihlal_risk_degerlendirmesi}"),
    ("h2", "4.4. KURULA, İLGİLİ KİŞİLERE VE GEREKTİĞİNDE ÜÇÜNCÜ KİŞİLERE BİLDİRİM YAPILMASI"),
    ("p", "VERİ SORUMLUSU, Kişisel Verileri Koruma Kurulu'nun Kişisel Veri İhlali Bildirim Usul ve Esaslarına İlişkin "
          "24.01.2019 tarihli ve 2019/10 sayılı Kararı uyarınca, kişisel veri ihlalinin öğrenildiği tarihten itibaren en geç "
          "72 saat içinde veri ihlalini KURUL'a bildirir."),
    ("p", "KURUL'a yapılacak bildirimde Kişisel Verileri Koruma Kurumu'nun internet sitesinde yayınlanmış olan Kişisel Veri "
          "İhlali Başvuru Formu kullanılır. Formda yer alan bilgilerin aynı anda sağlanmasının mümkün olmadığı hallerde, bu "
          "bilgiler gecikmeye mahal verilmeksizin aşamalı olarak sağlanabilir."),
    ("p", "Haklı bir gerekçe ile 72 saat içerisinde KURUL'a bildirim yapılamaması durumunda, yapılacak bildirimle birlikte "
          "gecikmenin nedenleri de KURUL'a açıklanır."),
    ("p", "Veri ihlalinden etkilenen kişilerin belirlenmesini müteakip ilgili kişilere de makul olan en kısa süre içerisinde "
          "ilgili kişinin iletişim adresine ulaşılabiliyorsa doğrudan, ulaşılamıyorsa VERİ SORUMLUSU'nun kendi internet "
          "sitesi üzerinden yayımlanması gibi uygun yöntemlerle bildirir."),
    ("p", "Kişisel Verileri Koruma Kurulunun 18.09.2019 tarihli ve 2019/271 sayılı Kararı gereğince; VERİ SORUMLUSU "
          "tarafından ilgili kişiye yapılacak olan ihlal bildirimi açık ve sade bir dille yapılır ve asgari olarak;"),
    ("madde", "İhlalin ne zaman gerçekleştiği,"),
    ("madde", "Kişisel veri kategorileri bazında (kişisel veri/özel nitelikli kişisel veri ayrımı yapılarak) hangi kişisel "
              "verilerin ihlalden etkilendiği,"),
    ("madde", "Kişisel veri ihlalinin muhtemel sonuçları,"),
    ("madde", "Veri ihlalinin olumsuz etkilerinin azaltılması için alınan veya alınması önerilen tedbirler,"),
    ("madde", "İlgili kişilerin veri ihlali ile ilgili bilgi almalarını sağlayacak irtibat kişilerinin isim ve iletişim "
              "detayları ya da veri sorumlusunun web sayfasının tam adresi, çağrı merkezi vb. iletişim yolları,"),
    ("p", "unsurlarına yer verilir."),
    ("p", "VERİ SORUMLUSU'nun kişisel veri işleme envanterine göre bir ihlalden etkilenebilecek veri kategorileri, "
          "bildirimde kullanılacak kişisel veri / özel nitelikli kişisel veri ayrımıyla aşağıda listelenmiştir:"),
    ("madde", "${ihlal_veri_kategorileri}"),
    ("p", "Veri İhlali Müdahale Ekibi tarafından veri ihlalinin niteliği ve kapsamı, ihlalin suç teşkil edip etmediği gibi "
          "hususlar değerlendirilerek diğer veri sorumluları, veri işleyenler, dış danışmanlar, adli makamlar ve bankalar "
          "gibi üçüncü kişilere de bildirim yapılabilir."),
    ("h2", "4.5. DEĞERLENDİRME VE GELİŞTİRME ÇALIŞMALARININ YÜRÜTÜLMESİ"),
    ("p", "Veri ihlalini takiben son aşamada Veri İhlali Müdahale Ekibi tarafından; muhtemel kişisel veri ihlallerinin "
          "etkilerini azaltmak için hangi adımların atılması gerektiği, bunun için ek bir idari ve/veya teknik tedbir "
          "alınmasının gerekip gerekmediği, kişisel veri ihlali nedeniyle herhangi bir politika veya prosedürde değişiklik "
          "ya da iyileştirme gerekip gerekmediği, tekrar bir veri ihlali ile karşı karşıya kalmamak için ek kaynak ve/veya "
          "alt yapıya ihtiyaç olup olmadığı, eğitim gerekliliğinin olup olmadığı gibi hususları içeren bir rapor hazırlanır "
          "ve Koordinatöre sunulur."),
    ("h1", "5. PROSEDÜRÜN GÜNCELLENMESİ"),
    ("p", "İşbu PROSEDÜR, kurumsal ya da mevzuattan kaynaklanan içerik değişiklik gereksinimlerine bakılmaksızın yılda bir "
          "kez gözden geçirilip gerektiğinde güncellenir."),
    ("p", "VERİ SORUMLUSU'nun ayrıca KANUN, Kişisel Verileri Koruma Kurulu kararları uyarınca ya da mevzuattaki gelişmeler "
          "doğrultusunda işbu PROSEDÜR metninde değişiklik yapma hakkını saklı tutar."),
    ("bos", ""),
    ("p", "Bu prosedür ${tarih} tarihinde ${kurum} tarafından onaylanmıştır."),
]


def _prototipler(d: docx.document.Document) -> dict[str, Paragraph]:
    # Kaynak sablondan bicim ornekleri: baslik (${KURUM}), h1, h2, govde, madde imli
    out: dict[str, Paragraph] = {}
    for p in d.paragraphs:
        metin = p.text.strip()
        if "baslik" not in out and metin == "${KURUM}":
            out["baslik"] = p
        elif "h1" not in out and p.style.name == "Heading 1" and metin.startswith("1."):
            out["h1"] = p
        elif "h2" not in out and p.style.name == "Heading 2":
            out["h2"] = p
        elif "p" not in out and p.style.name == "Normal" and len(metin) > 60 and p.runs:
            out["p"] = p
        elif "madde" not in out and p.style.name == "List Paragraph" and p._p.pPr is not None \
                and p._p.pPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numPr") is not None:
            out["madde"] = p
    eksik = {"baslik", "h1", "h2", "p", "madde"} - set(out)
    if eksik:
        raise SystemExit(f"Kaynak şablonda biçim örneği bulunamadı: {sorted(eksik)}")
    return out


def _yeni_paragraf(body, proto: Paragraph, metin: str) -> Paragraph:
    el = copy.deepcopy(proto._p)
    body.append(el)
    par = Paragraph(el, proto._parent)
    if par.runs:
        par.runs[0].text = metin
        for r in par.runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        par.add_run(metin)
    return par


def uret() -> Path:
    d = docx.Document(str(KAYNAK))
    proto = _prototipler(d)
    body = d.element.body
    sect = body.sectPr
    for el in list(body):
        body.remove(el)

    for tur, metin in ICERIK:
        if tur == "sayfa":
            p = _yeni_paragraf(body, proto["p"], "")
            p.add_run().add_break(WD_BREAK.PAGE)
        elif tur == "bos":
            _yeni_paragraf(body, proto["p"], "")
        elif tur in ("altbaslik", "orta"):
            p = _yeni_paragraf(body, proto["h1"] if tur == "altbaslik" else proto["p"], metin)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if tur == "orta":
                for r in p.runs:
                    r.font.bold = True
        else:
            _yeni_paragraf(body, proto[tur], metin)

    if sect is not None:
        body.append(sect)
    d.core_properties.title = "Veri İhlali Müdahale Prosedürü"
    d.save(str(HEDEF))
    return HEDEF


if __name__ == "__main__":
    yol = uret()
    print(f"Şablon yazıldı: {yol}")
    sys.exit(0)
