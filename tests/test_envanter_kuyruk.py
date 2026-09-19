# Envanter kuyruk isleri: Excel disa aktarim (POST /export -> islem -> dosya -> indir),
# Excel'den ice aktarim (yukleme -> kuyruk -> satirlar; department kapsami atlar),
# yeniden indeksleme (kuyruk), islem listesi/sahiplik, bildirim + denetim gunlugu.
import asyncio
import io
import os
import pathlib
import shutil
import tempfile
import unittest

try:
    import httpx
except ImportError:  # pragma: no cover
    httpx = None
try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None

from kvkk_rag.config import settings

_TMP = pathlib.Path(tempfile.mkdtemp(prefix="kvkk_envkuyruk_"))
settings.SQLITE_PATH = _TMP / "envkuyruk.db"
settings.INDEX_DIR = _TMP / "index"
os.environ["KVKK_ADMIN_EMAIL"] = "admin@test.local"
os.environ["KVKK_ADMIN_PASSWORD"] = "admin-sifre-123"

from kvkk_rag.api.app.Models.audit_log import AuditLog  # noqa: E402
from kvkk_rag.api.app.Models.envanter_islemi import EnvanterIslemi  # noqa: E402
from kvkk_rag.api.app.Models.notification import Notification  # noqa: E402
from kvkk_rag.api.app.Services import hashing  # noqa: E402
from kvkk_rag.api.app.Services.permission_cache import MemoryStore, cache  # noqa: E402
from kvkk_rag.api.app.Services.queue import Worker  # noqa: E402
from kvkk_rag.api.database import connection  # noqa: E402
from kvkk_rag.inventory import store as inv_store  # noqa: E402

hashing.ITERATIONS = 1000


def xlsx_olustur(satirlar: list[dict]) -> bytes:
    basliklar = ["Birim", "Faaliyet", "Veri Kategorisi", "Kişisel Veri", "Hukuki Sebep", "Saklama Süresi"]
    anahtar = ["birim", "faaliyet", "veri_kategorisi", "kisisel_veri", "hukuki_sebep", "saklama_suresi"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(basliklar)
    for s in satirlar:
        ws.append([s.get(k, "") for k in anahtar])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@unittest.skipIf(httpx is None or not hasattr(httpx, "ASGITransport") or openpyxl is None,
                 "httpx.ASGITransport ve openpyxl gerekli")
class TestEnvanterKuyruk(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from kvkk_rag.api import bootstrap

        settings.SQLITE_PATH = _TMP / "envkuyruk.db"
        settings.INDEX_DIR = _TMP / "index"
        connection.reset()
        cache.store = MemoryStore()
        cls.app = bootstrap.create_app(graphql=False, static=False, seed=False, worker=False)
        conn = connection.connect()
        for birim, veri in (("İK", "Ad Soyad"), ("İK", "TC Kimlik"), ("Pazarlama", "E-posta")):
            inv_store.create(conn, {"birim": birim, "faaliyet": "İşe alım", "kisisel_veri": veri,
                                    "veri_kategorisi": "Kimlik"}, kaynak="xlsx")
        conn.close()
        bootstrap.migrate_and_seed()

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(_TMP, ignore_errors=True)

    def client(self):
        return httpx.AsyncClient(transport=httpx.ASGITransport(app=self.app), base_url="http://test")

    async def login(self, c, email, password):
        r = await c.post("/api/auth/login", json={"email": email, "password": password})
        self.assertEqual(r.status_code, 200, r.text)
        d = r.json()
        return {"Authorization": f"Bearer {d['access_token']}"}, d

    def calistir(self):
        # Kuyrugu bosaltir (envanter degisimleri faaliyet belgesi islerini de kuyruga yazar)
        conn = connection.connect()
        try:
            n = 0
            while Worker.run_once(conn):
                n += 1
            return n > 0
        finally:
            conn.close()

    def test_disa_aktarim_kuyrugu(self):
        async def akis():
            async with self.client() as c:
                h, giris = await self.login(c, "admin@test.local", "admin-sifre-123")
                r = await c.post("/api/inventory/export", headers=h, json={"birim": "İK", "kurum": "TEST A.Ş."})
                self.assertEqual(r.status_code, 202, r.text)
                d = r.json()
                self.assertTrue(d["kuyrukta"])
                islem_id = d["islem"]["id"]
                self.assertEqual(d["islem"]["durum"], "kuyrukta")
                self.assertEqual(d["islem"]["tur"], "disa_aktar")

                r = await c.get(f"/api/inventory/islemler/{islem_id}/indir", headers=h)
                self.assertEqual(r.status_code, 409)

                self.assertTrue(self.calistir())
                r = await c.get(f"/api/inventory/islemler/{islem_id}", headers=h)
                i = r.json()["islem"]
                self.assertEqual(i["durum"], "tamamlandi", i.get("hata"))
                self.assertEqual(i["sonuc"]["satir"], 2)  # yalnizca İK
                self.assertTrue(i["indirilebilir"])

                r = await c.get(f"/api/inventory/islemler/{islem_id}/indir", headers=h)
                self.assertEqual(r.status_code, 200, r.text)
                self.assertIn("spreadsheetml", r.headers["content-type"])
                wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True)
                ws = wb[wb.sheetnames[0]]
                satirlar = [row for row in ws.iter_rows(values_only=True) if any(v is not None for v in row)]
                self.assertEqual(len(satirlar) - 1, 2)  # baslik + 2 satir

                # bildirim + gunluk
                r = await c.get("/api/notifications", headers=h)
                self.assertTrue(any(b["type"] == "inventory.export_ready" and b["data"]["islem_id"] == islem_id
                                    for b in r.json()["bildirimler"]))
                r = await c.get("/api/audit-logs", headers=h, params={"action": "inventory.export"})
                self.assertGreaterEqual(r.json()["toplam"], 1)

                # liste + sil
                r = await c.get("/api/inventory/islemler", headers=h)
                self.assertEqual(r.status_code, 200)
                self.assertTrue(any(x["id"] == islem_id for x in r.json()["islemler"]))
                r = await c.delete(f"/api/inventory/islemler/{islem_id}", headers=h)
                self.assertEqual(r.status_code, 200)

        asyncio.run(akis())

    def test_ice_aktarim_kuyrugu_ve_kapsam(self):
        async def akis():
            async with self.client() as c:
                h, _ = await self.login(c, "admin@test.local", "admin-sifre-123")
                # department kapsamli kullanici: yalnizca kendi biriminin satirlari eklenir
                r = await c.get("/api/departments", headers=h)
                ik = next(d["id"] for d in r.json()["departmanlar"] if d["name"] == "İK")
                r = await c.post("/api/users", headers=h, json={"name": "Ayşe", "email": "ayse@test.local",
                                                                 "password": "ayse-sifre-123",
                                                                 "roles": ["Birim Sorumlusu"], "departments": [ik]})
                self.assertEqual(r.status_code, 200, r.text)
                h2, _ = await self.login(c, "ayse@test.local", "ayse-sifre-123")

                dosya = xlsx_olustur([
                    {"birim": "İK", "faaliyet": "Bordro", "veri_kategorisi": "Finans", "kisisel_veri": "IBAN",
                     "hukuki_sebep": "Kanunlarda Açıkça Öngörülmesi", "saklama_suresi": "10 yıl"},
                    {"birim": "Pazarlama", "faaliyet": "Kampanya", "veri_kategorisi": "İletişim", "kisisel_veri": "Telefon"},
                ])
                # yanlis uzanti
                r = await c.post("/api/inventory/import", headers=h2, files={"dosya": ("veri.csv", b"a,b", "text/csv")})
                self.assertEqual(r.status_code, 422)
                # degistir modu: tam kapsamli silme yetkisi gerekir
                r = await c.post("/api/inventory/import", headers=h2, data={"mod": "degistir"},
                                 files={"dosya": ("veri.xlsx", dosya, "application/octet-stream")})
                self.assertEqual(r.status_code, 403)
                # ekle modu
                r = await c.post("/api/inventory/import", headers=h2, data={"mod": "ekle"},
                                 files={"dosya": ("yeni_veri.xlsx", dosya, "application/octet-stream")})
                self.assertEqual(r.status_code, 202, r.text)
                islem_id = r.json()["islem"]["id"]
                self.assertEqual(r.json()["islem"]["ad"], "yeni_veri.xlsx")

                self.assertTrue(self.calistir())
                r = await c.get(f"/api/inventory/islemler/{islem_id}", headers=h2)
                i = r.json()["islem"]
                self.assertEqual(i["durum"], "tamamlandi", i.get("hata"))
                self.assertEqual(i["sonuc"]["eklenen"], 1)
                self.assertEqual(i["sonuc"]["atlanan"], 1)  # Pazarlama satiri kapsam disi

                r = await c.get("/api/inventory", headers=h2, params={"arama": "IBAN"})
                self.assertEqual(r.json()["filtrelenmis"], 1)
                # eklenen satir kullanicinin kendi kaydi (created_by) ve envanter_log'da
                conn = connection.connect()
                try:
                    self.assertTrue(AuditLog.query(conn).where("action", "inventory.import").exists())
                    n = Notification.query(conn).where("type", "inventory.import_done").first()
                    self.assertIsNotNone(n)
                    self.assertIn("1 satır eklendi", n.body)
                    self.assertIn("1 satır kapsam dışı", n.body)
                finally:
                    conn.close()

                # baskasinin islemi gorunmez
                r = await c.get(f"/api/inventory/islemler/{islem_id}", headers=h)   # admin (super) gorebilir
                self.assertEqual(r.status_code, 200)
                r = await c.post("/api/users", headers=h, json={"name": "Can", "email": "can@test.local",
                                                                 "password": "can-sifre-123", "roles": ["Denetçi"]})
                h3, _ = await self.login(c, "can@test.local", "can-sifre-123")
                r = await c.get(f"/api/inventory/islemler/{islem_id}", headers=h3)
                self.assertEqual(r.status_code, 403)

                # admin degistir modu: envanter bastan yazilir
                r = await c.post("/api/inventory/import", headers=h, data={"mod": "degistir"},
                                 files={"dosya": ("tam.xlsx", dosya, "application/octet-stream")})
                self.assertEqual(r.status_code, 202, r.text)
                self.assertTrue(self.calistir())
                r = await c.get("/api/inventory", headers=h)
                self.assertEqual(r.json()["toplam"], 2)

        asyncio.run(akis())

    def test_yeniden_indeksleme_kuyrugu(self):
        async def akis():
            async with self.client() as c:
                h, _ = await self.login(c, "admin@test.local", "admin-sifre-123")
                r = await c.post("/api/inventory/reindex", headers=h)
                self.assertEqual(r.status_code, 202, r.text)
                islem_id = r.json()["islem"]["id"]
                self.assertEqual(r.json()["islem"]["tur"], "yeniden_indeksle")
                self.calistir()  # lancedb yoksa is 'hata' ile biter; her iki durumda kayit kapanir ve bildirim duser
                r = await c.get(f"/api/inventory/islemler/{islem_id}", headers=h)
                self.assertIn(r.json()["islem"]["durum"], ("tamamlandi", "hata"))
                r = await c.get("/api/notifications", headers=h)
                self.assertTrue(any(b["data"].get("islem_id") == islem_id for b in r.json()["bildirimler"]))

        asyncio.run(akis())


if __name__ == "__main__":
    unittest.main()
