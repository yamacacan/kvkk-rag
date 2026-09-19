# Faaliyet belgeleri (envanter degisince otomatik yenilenen Aydinlatma Metni + Acik Riza Beyani)
# ve acik riza kayitlari (consents modulu: izin, kapsam, dogrulama, denetim gunlugu).
import asyncio
import os
import pathlib
import shutil
import tempfile
import unittest

try:
    import httpx
except ImportError:  # pragma: no cover
    httpx = None

from kvkk_rag.config import settings

_TMP = pathlib.Path(tempfile.mkdtemp(prefix="kvkk_faaliyet_"))
settings.SQLITE_PATH = _TMP / "faaliyet.db"
settings.INDEX_DIR = _TMP / "index"
os.environ["KVKK_ADMIN_EMAIL"] = "admin@test.local"
os.environ["KVKK_ADMIN_PASSWORD"] = "admin-sifre-123"

from kvkk_rag.api.app.Models.audit_log import AuditLog  # noqa: E402
from kvkk_rag.api.app.Models.faaliyet_belgesi import FaaliyetBelgesi  # noqa: E402
from kvkk_rag.api.app.Models.job import Job as JobRow  # noqa: E402
from kvkk_rag.api.app.Services import hashing  # noqa: E402
from kvkk_rag.api.app.Services.permission_cache import MemoryStore, cache  # noqa: E402
from kvkk_rag.api.app.Services.queue import Worker  # noqa: E402
from kvkk_rag.api.database import connection  # noqa: E402
from kvkk_rag.compliance import store as comp_store  # noqa: E402
from kvkk_rag.inventory import store as inv_store  # noqa: E402

hashing.ITERATIONS = 1000
TC_GECERLI = "10000000146"  # kontrol basamaklari gecerli ornek numara


@unittest.skipIf(httpx is None or not hasattr(httpx, "ASGITransport"), "httpx.ASGITransport gerekli")
class TestFaaliyetBelgeleriVeAcikRiza(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from kvkk_rag.api import bootstrap

        settings.SQLITE_PATH = _TMP / "faaliyet.db"
        settings.INDEX_DIR = _TMP / "index"
        connection.reset()
        cache.store = MemoryStore()
        cls.app = bootstrap.create_app(graphql=False, static=False, seed=False, worker=False)
        conn = connection.connect()
        conn.executescript(comp_store.SCHEMA)
        comp_store.save(conn, {"kurum": "TEST A.Ş.", "adres": "Test Mah. No:1", "web_adres": "https://test.local"})
        for birim, faaliyet, veri in (("İK", "İşe alım", "Özgeçmiş"), ("İK", "İşe alım", "Ad Soyad"),
                                      ("Pazarlama", "Kampanya", "E-posta")):
            inv_store.create(conn, {"birim": birim, "faaliyet": faaliyet, "kisisel_veri": veri, "veri_kategorisi": "Kimlik",
                                    "isleme_amaci": f"{faaliyet} süreçlerinin yürütülmesi", "alici_grubu": "SGK",
                                    "hukuki_sebep": "İlgili Kişinin Açık Rızasının Varlığı"}, kaynak="xlsx")
        conn.close()
        bootstrap.migrate_and_seed()

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(_TMP, ignore_errors=True)

    def client(self):
        return httpx.AsyncClient(transport=httpx.ASGITransport(app=self.app), base_url="http://test")

    async def login(self, c, email, password):
        r = await c.post("/api/auth/login", json={"email": email, "password": password})
        self.assertEqual(r.status_code, 200, r.text)
        d = r.json()
        return {"Authorization": f"Bearer {d['access_token']}"}, d

    def kuyrugu_bosalt(self):
        conn = connection.connect()
        try:
            n = 0
            while Worker.run_once(conn):
                n += 1
            return n
        finally:
            conn.close()

    def test_faaliyet_belgeleri_otomatik_yenilenir(self):
        async def akis():
            async with self.client() as c:
                h, _ = await self.login(c, "admin@test.local", "admin-sifre-123")
                # ilk listeleme: eksik belgeler kuyruga alinir
                r = await c.get("/api/faaliyet-belgeleri", headers=h)
                self.assertEqual(r.status_code, 200, r.text)
                d = r.json()
                self.assertEqual(d["ozet"]["faaliyet"], 2)
                self.assertEqual(set(d["sablonlar"]), {"aydinlatma", "acik_riza"})
                self.assertGreaterEqual(d["ozet"]["kuyrukta"], 1)
                self.assertGreaterEqual(self.kuyrugu_bosalt(), 4)

                r = await c.get("/api/faaliyet-belgeleri", headers=h)
                d = r.json()
                self.assertEqual(d["ozet"]["guncel"], 4, d["ozet"])
                ise_alim = next(f for f in d["faaliyetler"] if f["faaliyet"] == "İşe alım")
                self.assertEqual(ise_alim["birim"], "İK")
                self.assertEqual(ise_alim["satir"], 2)
                beyan = ise_alim["belgeler"]["acik_riza"]
                self.assertEqual(beyan["durum"], "guncel")
                self.assertTrue(beyan["indirilebilir"])
                r = await c.get(f"/api/faaliyet-belgeleri/{beyan['id']}/indir", headers=h)
                self.assertEqual(r.status_code, 200)
                self.assertIn("wordprocessingml", r.headers["content-type"])
                import docx, io
                metin = "\n".join(p.text for p in docx.Document(io.BytesIO(r.content)).paragraphs)
                self.assertIn("AÇIK RIZA BEYANI (İK İŞLEMLERİ)", metin)
                self.assertIn("Özgeçmiş", metin)
                self.assertIn("Ad Soyad", metin)
                self.assertIn("SGK", metin)

                # parmak izi degismedi -> yeniden uretim atlanir
                conn = connection.connect()
                from kvkk_rag.api.app.Services.faaliyet_belge_service import FaaliyetBelgeService
                try:
                    self.assertEqual(FaaliyetBelgeService.uret(conn, "İşe alım", "acik_riza")[0], "guncel")
                finally:
                    conn.close()

                # envanter degisti (yeni kisisel veri) -> belge eski + kuyruk -> yeniden uretim
                r = await c.post("/api/inventory", headers=h, json={"birim": "İK", "faaliyet": "İşe alım",
                                                                   "kisisel_veri": "Adli Sicil Kaydı", "veri_kategorisi": "Ceza Mahkumiyeti ve Güvenlik Tedbirleri"})
                self.assertEqual(r.status_code, 200, r.text)
                conn = connection.connect()
                try:
                    b = FaaliyetBelgesi.find_for(conn, "İşe alım", "acik_riza")
                    self.assertEqual(b.durum, "eski")
                    self.assertTrue(JobRow.query(conn).where("status", "queued").exists())
                finally:
                    conn.close()
                self.kuyrugu_bosalt()
                r = await c.get(f"/api/faaliyet-belgeleri/{beyan['id']}/indir", headers=h)
                metin = "\n".join(p.text for p in docx.Document(io.BytesIO(r.content)).paragraphs)
                self.assertIn("Adli Sicil Kaydı", metin)
                conn = connection.connect()
                try:
                    self.assertTrue(AuditLog.query(conn).where("action", "documents.auto_refresh").exists())
                    # kurum profili degisince hepsi yenilenir (toplu is)
                finally:
                    conn.close()
                r = await c.put("/api/profile", headers=h, json={"kurum": "YENİ KURUM A.Ş.", "adres": "Test Mah. No:1", "web_adres": "https://test.local"})
                self.assertEqual(r.status_code, 200, r.text)
                self.kuyrugu_bosalt()
                r = await c.get(f"/api/faaliyet-belgeleri/{beyan['id']}/indir", headers=h)
                metin = "\n".join(p.text for p in docx.Document(io.BytesIO(r.content)).paragraphs)
                self.assertIn("YENİ KURUM A.Ş.", metin)

                # elle yenile
                r = await c.post("/api/faaliyet-belgeleri/yenile", headers=h, params={"faaliyet": "Kampanya"})
                self.assertEqual(r.status_code, 200, r.text)

                # department kapsamli kullanici yalnizca kendi faaliyetlerini gorur
                r = await c.get("/api/departments", headers=h)
                ik = next(d["id"] for d in r.json()["departmanlar"] if d["name"] == "İK")
                r = await c.post("/api/users", headers=h, json={"name": "Ayşe", "email": "ayse@test.local", "password": "ayse-sifre-123",
                                                                 "roles": ["Birim Sorumlusu"], "departments": [ik]})
                self.assertEqual(r.status_code, 200, r.text)
                h2, _ = await self.login(c, "ayse@test.local", "ayse-sifre-123")
                r = await c.get("/api/faaliyet-belgeleri", headers=h2)
                self.assertEqual([f["faaliyet"] for f in r.json()["faaliyetler"]], ["İşe alım"])
                kampanya = next(f for f in d["faaliyetler"] if f["faaliyet"] == "Kampanya")
                r = await c.get(f"/api/faaliyet-belgeleri/{kampanya['belgeler']['aydinlatma']['id']}/indir", headers=h2)
                self.assertEqual(r.status_code, 403)

        asyncio.run(akis())

    def test_acik_riza_kayitlari(self):
        async def akis():
            async with self.client() as c:
                h, _ = await self.login(c, "admin@test.local", "admin-sifre-123")
                r = await c.get("/api/consents/faaliyetler", headers=h)
                self.assertEqual(r.status_code, 200, r.text)
                self.assertEqual({f["faaliyet"] for f in r.json()["faaliyetler"]}, {"İşe alım", "Kampanya"})

                # dogrulama: gecersiz TC, eksik geri cekme tarihi
                r = await c.post("/api/consents", headers=h, json={"faaliyet": "İşe alım", "tc_kimlik": "12345678901",
                                                                   "ad": "Ali", "soyad": "Veli", "durum": "onaylandi"})
                self.assertEqual(r.status_code, 422, r.text)
                self.assertIn("Kimlik", r.json()["detail"])
                r = await c.post("/api/consents", headers=h, json={"faaliyet": "İşe alım", "tc_kimlik": TC_GECERLI,
                                                                   "ad": "Ali", "soyad": "Veli", "durum": "geri_cekildi"})
                self.assertEqual(r.status_code, 422, r.text)
                self.assertIn("geri çekme tarihi", r.json()["detail"])

                # olustur
                r = await c.post("/api/consents", headers=h, json={"faaliyet": "İşe alım", "tc_kimlik": TC_GECERLI,
                                                                   "ad": "Ali", "soyad": "Veli", "durum": "onaylandi",
                                                                   "onay_yontemi": "islak_imza", "notlar": "Form imzalandı"})
                self.assertEqual(r.status_code, 200, r.text)
                k = r.json()["kayit"]
                self.assertEqual(k["birim"], "İK")
                self.assertTrue(k["onay_tarihi"])
                self.assertEqual(k["tc_kimlik"], TC_GECERLI)  # tekil kayit: maskesiz
                cid = k["id"]

                # liste: maskeli TC, ozet
                r = await c.get("/api/consents", headers=h)
                d = r.json()
                self.assertEqual(d["toplam"], 1)
                self.assertEqual(d["kayitlar"][0]["tc_kimlik"], "100*****146")
                self.assertEqual(d["ozet"]["onaylandi"], 1)
                self.assertEqual(d["kayitlar"][0]["durum_adi"], "Onaylandı")

                # geri cek
                r = await c.patch(f"/api/consents/{cid}", headers=h, json={"faaliyet": "İşe alım", "tc_kimlik": TC_GECERLI,
                                                                            "ad": "Ali", "soyad": "Veli", "durum": "geri_cekildi",
                                                                            "geri_cekme_tarihi": "2026-09-19T10:00", "onay_yontemi": "eposta"})
                self.assertEqual(r.status_code, 200, r.text)
                self.assertEqual(r.json()["kayit"]["durum"], "geri_cekildi")
                r = await c.get("/api/consents", headers=h, params={"durum": "geri_cekildi"})
                self.assertEqual(r.json()["toplam"], 1)

                # department kapsamli kullanici: kendi birimi disinda kayit acamaz, digerlerini goremez
                r = await c.get("/api/departments", headers=h)
                pz = next(d["id"] for d in r.json()["departmanlar"] if d["name"] == "Pazarlama")
                r = await c.post("/api/users", headers=h, json={"name": "Can", "email": "can@test.local", "password": "can-sifre-123",
                                                                 "roles": ["Birim Sorumlusu"], "departments": [pz]})
                self.assertEqual(r.status_code, 200, r.text)
                h2, _ = await self.login(c, "can@test.local", "can-sifre-123")
                r = await c.get("/api/consents", headers=h2)
                self.assertEqual(r.json()["toplam"], 0)
                self.assertEqual(r.json()["kapsam"], "department")
                r = await c.post("/api/consents", headers=h2, json={"faaliyet": "İşe alım", "tc_kimlik": TC_GECERLI,
                                                                    "ad": "Zeynep", "soyad": "Kaya", "durum": "onaylandi"})
                self.assertEqual(r.status_code, 403)
                r = await c.post("/api/consents", headers=h2, json={"faaliyet": "Kampanya", "tc_kimlik": TC_GECERLI,
                                                                    "ad": "Zeynep", "soyad": "Kaya", "durum": "onaylanmadi"})
                self.assertEqual(r.status_code, 200, r.text)
                r = await c.get(f"/api/consents/{cid}", headers=h2)
                self.assertEqual(r.status_code, 403)
                r = await c.delete(f"/api/consents/{cid}", headers=h2)
                self.assertEqual(r.status_code, 403)

                # sil + gunluk
                r = await c.delete(f"/api/consents/{cid}", headers=h)
                self.assertEqual(r.status_code, 200)
                r = await c.get("/api/audit-logs", headers=h, params={"action": "consents"})
                eylemler = {k["action"] for k in r.json()["kayitlar"]}
                self.assertEqual(eylemler, {"consents.create", "consents.update", "consents.delete"})
                # gunlukte TC maskeli
                kayit = next(k for k in r.json()["kayitlar"] if k["action"] == "consents.create")
                self.assertEqual(kayit["after"]["tc_kimlik"], "100*****146")

        asyncio.run(akis())


if __name__ == "__main__":
    unittest.main()
