# Veri İhlali Müdahale Prosedürü şablonu: kurum/tarih placeholder'ları profilden,
# birim ve veri kategorisi listeleri envanterden, senaryo/risk bölümleri LLM'den
# (LLM yoksa deterministik yedek) dolar; belgede ${...} kalıntısı kalmaz.
import io
import unittest

try:
    import docx
except ImportError:  # pragma: no cover
    docx = None

from kvkk_rag.compliance import generate as comp_generate
from kvkk_rag.compliance import profile as comp_profile
from kvkk_rag.compliance import sections
from kvkk_rag.config import settings
from kvkk_rag.inventory.loader import InventoryRow

ROWS = [
    InventoryRow(satir_no=1, birim="İNSAN KAYNAKLARI", faaliyet="İşe alım", veri_kategorisi="Kimlik",
                 kisisel_veri="Ad Soyad", kayit_ortami="Fiziksel arşiv", alici_grubu="SGK",
                 hukuki_sebep="Kanunlarda Açıkça Öngörülmesi", saklama_suresi="5 yıl", imha_yontemi="silme",
                 teknik_tedbir="Şifreleme", idari_tedbir="Eğitim"),
    InventoryRow(satir_no=2, birim="SAĞLIK BİRİMİ", faaliyet="Muayene", veri_kategorisi="Sağlık Bilgileri",
                 kisisel_veri="Tanı", kayit_ortami="HBYS", alici_grubu="Hastane",
                 hukuki_sebep="Kanunlarda Açıkça Öngörülmesi", saklama_suresi="15 yıl",
                 imha_yontemi="anonimleştirme", teknik_tedbir="Erişim logları", idari_tedbir="Gizlilik taahhütnamesi"),
]


class SahteLLM:
    def complete(self, messages, **kw):
        return "Yapay zekâ ile yazılmış, yalnızca verilen olgulara dayanan yeterince uzun bir kurumsal paragraf. " * 4


@unittest.skipIf(docx is None or not (settings.TEMPLATES_DIR / comp_generate.SABLONLAR["veri_ihlali"]).exists(),
                 "python-docx ve şablon dosyası gerekli")
class TestVeriIhlaliSablonu(unittest.TestCase):
    def _profil(self):
        p = comp_profile.from_inventory(ROWS, kurum="ÖRNEK TEKNOLOJİ A.Ş.")
        p.belge_tarihi = "19.09.2026"
        return p

    def test_llmsiz_uretim(self):
        icerik, kalan, ad, kaynak = comp_generate.uret("veri_ihlali", self._profil(), ROWS, None)
        self.assertEqual(kalan, [])
        self.assertEqual(ad, "Veri_İhlali_Müdahale_Prosedürü.docx")
        self.assertEqual(kaynak["ihlal_birimleri"], "veritabani")
        self.assertEqual(kaynak["ihlal_senaryolari"], "yedek")
        metin = "\n".join(p.text for p in docx.Document(io.BytesIO(icerik)).paragraphs)
        self.assertNotIn("${", metin)
        self.assertNotIn("4Dimension", metin)
        self.assertIn("ÖRNEK TEKNOLOJİ A.Ş.", metin)
        self.assertIn("Tarih: 19.09.2026", metin)
        self.assertIn("Bu prosedür 19.09.2026 tarihinde ÖRNEK TEKNOLOJİ A.Ş. tarafından onaylanmıştır.", metin)
        # envanterden gelen listeler madde imli paragraflara acilir
        self.assertIn("SAĞLIK BİRİMİ", metin)
        self.assertIn("Sağlık Bilgileri — özel nitelikli kişisel veri (KANUN m.6)", metin)
        self.assertIn("Kimlik — kişisel veri", metin)
        self.assertIn("72 saat", metin)

    def test_llm_ile_uretim(self):
        _, kalan, _, kaynak = comp_generate.uret("veri_ihlali", self._profil(), ROWS, SahteLLM())
        self.assertEqual(kalan, [])
        self.assertEqual(kaynak["ihlal_senaryolari"], "yapay_zeka")
        self.assertEqual(kaynak["ihlal_risk_degerlendirmesi"], "yapay_zeka")

    def test_sablon_bilgisi(self):
        bilgi = {s["anahtar"]: s for s in comp_generate.sablon_bilgisi()}["veri_ihlali"]
        self.assertTrue(bilgi["mevcut"])
        self.assertEqual(set(bilgi["uretilen_bolumler"]),
                         {"ihlal_birimleri", "ihlal_veri_kategorileri", "ihlal_senaryolari", "ihlal_risk_degerlendirmesi"})
        self.assertTrue(set(sections.YAPAY_ZEKA) >= {"ihlal_senaryolari", "ihlal_risk_degerlendirmesi"})


@unittest.skipIf(docx is None or not (settings.TEMPLATES_DIR / comp_generate.SABLONLAR["basvuru_formu"]).exists(),
                 "python-docx ve şablon dosyası gerekli")
class TestBasvuruFormuSablonu(unittest.TestCase):
    # Bos form: kurum (baslikta buyuk harf), adres, tarih dolar; ilgili kisi alanlari bos kalir
    def test_uretim(self):
        p = comp_profile.from_inventory(ROWS, kurum="Örnek Teknoloji A.Ş.", adres="Örnek Mah. No:1 Ankara")
        p.belge_tarihi = "19.09.2026"
        icerik, kalan, ad, kaynak = comp_generate.uret("basvuru_formu", p, ROWS, None)
        self.assertEqual(kalan, [])
        self.assertEqual(ad, "Kişisel_Veri_Sahibi_Başvuru_Formu.docx")
        d = docx.Document(io.BytesIO(icerik))
        metin = "\n".join(x.text for x in d.paragraphs)
        hucreler = [c.text for t in d.tables for r in t.rows for c in r.cells]
        self.assertNotIn("${", metin + "".join(hucreler))
        self.assertEqual(d.paragraphs[1].text, "ÖRNEK TEKNOLOJİ A.Ş.")  # ${KURUM} buyuk harf
        self.assertIn("Bu form 19.09.2026 tarihinde Örnek Teknoloji A.Ş. tarafından hazırlanmıştır.", metin)
        self.assertEqual(len(d.tables), 6)
        self.assertIn("Örnek Teknoloji A.Ş.\nÖrnek Mah. No:1 Ankara", hucreler)  # basvuru adresi hucresi
        self.assertIn("☐  Vatandaş", metin)
        self.assertIn("Adı Soyadı:", hucreler)
        self.assertIn("6698 sayılı Kişisel Verilerin Korunması Kanunu", hucreler)


if __name__ == "__main__":
    unittest.main()


@unittest.skipIf(docx is None or not (settings.TEMPLATES_DIR / comp_generate.SABLONLAR["genel_aydinlatma"]).exists(),
                 "python-docx ve şablon dosyası gerekli")
class TestGenelAydinlatmaSablonu(unittest.TestCase):
    # Tek sayfalik genel bildirim: kurum/adres/web/cagri merkezi profilden, QR kod web adresinden
    def _profil(self, web="https://www.ornekteknoloji.com.tr"):
        p = comp_profile.from_inventory(ROWS, kurum="ÖRNEK TEKNOLOJİ A.Ş.", adres="Örnek Mah. No:1 Ankara", web_adres=web)
        p.cagri_merkezi = "0850 000 00 00"
        return p

    def test_qr_kodlu_uretim(self):
        icerik, kalan, ad, kaynak = comp_generate.uret("genel_aydinlatma", self._profil(), ROWS, None)
        self.assertEqual(kalan, [])
        self.assertEqual(kaynak, {})
        d = docx.Document(io.BytesIO(icerik))
        metin = "\n".join(p.text for p in d.paragraphs)
        self.assertNotIn("${", metin)
        self.assertIn("Veri Sorumlusu ÖRNEK TEKNOLOJİ A.Ş. olarak", metin)
        self.assertIn("Çağrı Merkezi: 0850 000 00 00", metin)
        self.assertIn("Örnek Mah. No:1 Ankara", metin)
        self.assertEqual(metin.count("https://www.ornekteknoloji.com.tr"), 2)
        try:
            import segno  # noqa: F401
            self.assertEqual(len(d.inline_shapes), 1)  # QR resmi
        except ImportError:
            self.assertEqual(len(d.inline_shapes), 0)

    def test_web_adresi_yoksa_qr_bos(self):
        icerik, kalan, _, _ = comp_generate.uret("genel_aydinlatma", self._profil(web=""), ROWS, None)
        self.assertEqual(kalan, [])
        self.assertEqual(len(docx.Document(io.BytesIO(icerik)).inline_shapes), 0)

    def test_sablon_bilgisi_qr_alan_degil(self):
        bilgi = {s["anahtar"]: s for s in comp_generate.sablon_bilgisi()}["genel_aydinlatma"]
        self.assertNotIn("qr_kod", bilgi["alanlar"])
        self.assertIn("cagri_merkezi", bilgi["alanlar"])
